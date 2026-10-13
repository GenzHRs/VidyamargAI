"""Create reproducible synthetic fixtures with approved libraries and stdlib."""
from pathlib import Path
import zlib
from pypdf import PdfWriter, PdfReader
from pypdf.generic import DictionaryObject,NameObject,NumberObject,DecodedStreamObject,ArrayObject
from pdf2image import convert_from_path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT=Path(__file__).resolve().parents[1]/'tests/fixtures'


def pdf(path, pages):
    writer=PdfWriter()
    for commands in pages:
        page=writer.add_blank_page(612,792)
        regular=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Courier' if path.name=='scan_source.pdf' else '/Helvetica')})
        bold=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Courier-Bold' if path.name=='scan_source.pdf' else '/Helvetica-Bold')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(regular),NameObject('/F2'):writer._add_object(bold)})})
        stream=DecodedStreamObject();stream.set_data('\n'.join(commands).encode('latin-1'));page[NameObject('/Contents')]=writer._add_object(stream)
    writer.write(path)


def text(x,y,value,size=11,bold=False):
    escaped=value.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
    return f'BT /F{2 if bold else 1} {size} Tf 1 0 0 1 {x} {y} Tm ({escaped}) Tj ET'


def main():
    ROOT.mkdir(parents=True,exist_ok=True)
    sidebar=[text(35,750,'CONTACT',12,True),text(35,730,'jane@example.com'),text(35,710,'Seattle, WA'),text(35,675,'SKILLS',12,True),text(35,655,'Python SQL'),text(35,635,'Docker Linux'),text(220,750,'JANE DOE',24,True),text(220,722,'Software Engineer'),text(220,682,'EXPERIENCE',12,True),text(220,660,'Engineer | Example Systems',11,True),text(220,642,'2021 - Present'),text(220,620,'Built reliable services for customers.'),text(220,600,'Led the team through a cloud migration.')]
    pdf(ROOT/'two_column_sidebar.pdf',[sidebar])
    hero=[text(35,750,'JANE DOE',24,True),text(35,727,'Software Engineer'),text(35,707,'Seattle, WA | jane@example.com'),text(35,677,'SUMMARY',12,True),text(35,657,'Experienced software engineer building reliable systems for customers.'),text(35,607,'EXPERIENCE',12,True),text(35,587,'Engineer | Example Systems',11,True),text(35,567,'2021 - Present'),text(35,547,'Delivered reliable services.'),text(365,607,'SKILLS',12,True),text(365,587,'Python SQL Docker'),text(365,567,'Linux Kubernetes'),text(365,547,'EDUCATION',12,True),text(365,527,'B.S. in Computer Science')]
    pdf(ROOT/'horizontal_hero.pdf',[hero])
    dark=['0.05 0.10 0.20 rg 20 630 570 145 re f','1 1 1 rg',text(35,750,'JANE DOE',24,True),text(35,727,'Software Engineer'),text(35,707,'Seattle, WA | jane@example.com'),'0 0 0 rg']+hero[3:]
    pdf(ROOT/'dark_background.pdf',[dark])
    grid=[text(35,750,'JANE DOE',24,True),text(35,727,'jane@example.com | Seattle, WA'),text(35,687,'EXPERIENCE',12,True),'0 0 0 RG 0.5 w']
    grid += [f'{x} 500 m {x} 660 l S' for x in [35,155,325,577]]+[f'35 {y} m 577 {y} l S' for y in [500,570,660]]
    grid += [text(42,641,'2021 - Present'),text(163,641,'Senior Engineer',11,True),text(163,623,'Example Systems'),text(333,641,'Built reliable software services.'),text(333,623,'Managed production cloud systems.'),text(42,550,'2018 - 2021'),text(163,550,'Engineer',11,True),text(163,532,'Other Company'),text(333,550,'Improved customer service tools.'),text(333,532,'Maintained core platform systems.')]
    pdf(ROOT/'heavy_table.pdf',[grid])
    page1=[text(35,765,'JANE DOE',18,True),text(35,692,'EDUCATION',12,True),text(35,672,'Ph.D. in Computer Science'),text(35,652,'Example University 2018 - 2023'),text(35,612,'PUBLICATIONS',12,True),text(35,592,'Doe, J. (2024). Reliable systems and reproducible scientific experiments.'),text(35,572,'Journal of Software Research, volume 10, pages 120 - 130.'),text(35,35,'Confidential resume')]
    page2=[text(35,765,'JANE DOE',18,True),text(35,692,'EXPERIENCE',12,True),text(35,672,'Research Fellow | Example University',11,True),text(35,652,'2023 - Present'),text(35,612,'Developed methods for reliable systems and scientific experiments.'),text(35,592,'Collaborated with researchers across departments to publish findings.'),text(35,572,'Mentored students in reproducible software engineering and analysis.'),text(35,35,'Confidential resume')]
    pdf(ROOT/'academic_repeated_margins.pdf',[page1,page2])
    transcript=['JANE DOE','Software Engineer','jane@example.com','Seattle, WA','SUMMARY','Experienced software engineer building reliable services.','SKILLS','Python SQL Docker Linux','EXPERIENCE','Engineer at Example Systems','2021 - Present','Built reliable services for customers.','Led a team of engineers on cloud migrations.']
    pdf(ROOT/'scan_source.pdf',[[text(40,750-i*38,value,24 if i==0 else 14,i==0) for i,value in enumerate(transcript)]])
    image=convert_from_path(str(ROOT/'scan_source.pdf'),dpi=200,thread_count=1)[0].convert('RGB')
    writer=PdfWriter();page=writer.add_blank_page(612,792)
    obj=DecodedStreamObject();obj.set_data(image.tobytes());obj.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(image.width),NameObject('/Height'):NumberObject(image.height),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im0'):writer._add_object(obj.flate_encode())})})
    stream=DecodedStreamObject();stream.set_data(b'q 612 0 0 792 0 0 cm /Im0 Do Q');page[NameObject('/Contents')]=writer._add_object(stream);writer.write(ROOT/'scanned.pdf')
    (ROOT/'scanned_gold.txt').write_text('\n'.join(transcript))
    doc=Document();doc.sections[0].header.paragraphs[0].text='JANE DOE';doc.add_paragraph('Software Engineer')
    p=doc.add_paragraph();p.add_run('Contact: ')
    hyperlink=OxmlElement('w:hyperlink');run=OxmlElement('w:r');t=OxmlElement('w:t');t.text='jane@example.com';run.append(t);hyperlink.append(run);p._p.append(hyperlink)
    p=doc.add_paragraph();r=p.add_run('Seattle, WA');r.add_break();r.add_text('Phone: +1 (415) 555-0199');r.add_tab();r.add_text('Available')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).merge(table.cell(0,1)).text='Merged section heading';table.cell(1,0).text='2021 - Present';table.cell(1,1).text='Engineer | Example Systems'
    nested=table.cell(1,1).add_table(rows=1,cols=1);nested.cell(0,0).text='Nested content once'
    box=OxmlElement('w:txbxContent');p=OxmlElement('w:p');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Textbox content once';r.append(t);p.append(r);box.append(p);doc.add_paragraph()._p.append(box)
    sdt=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent');p=OxmlElement('w:p');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='Content control text';r.append(t);p.append(r);content.append(p);sdt.append(content);doc._element.body.insert(0,sdt)
    doc.add_paragraph('After the table');doc.save(ROOT/'docx_edge_cases.docx')

if __name__=='__main__':main()
